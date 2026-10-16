"""
Script para importar el catálogo de tiendas desde Hoja2 del Excel a la BD.
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

import openpyxl
from app import app
from database import db, Tienda


def importar_tiendas():
    filepath = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "datos.xlsx")
    wb = openpyxl.load_workbook(filepath, data_only=True)

    if "Hoja2" not in wb.sheetnames:
        print("ERROR: No se encontró Hoja2 en el Excel")
        sys.exit(1)

    ws = wb["Hoja2"]
    rows = list(ws.iter_rows(values_only=True))

    # Primera fila = headers
    headers = [str(h).strip() if h else "" for h in rows[0]]
    print(f"Columnas en Hoja2: {headers}")

    # Mapear columnas
    col_map = {}
    for i, h in enumerate(headers):
        col_map[h.upper()] = i

    col_tienda = col_map.get("TIENDA")
    col_co = col_map.get("CO")
    col_ciudad = col_map.get("CIUDAD")
    col_concepto = col_map.get("CONCEPTO")
    col_correo = col_map.get("CORREO TIENDA")
    col_direccion = col_map.get("DIRECCION TIENDA")

    with app.app_context():
        db.create_all()

        importados = 0
        actualizados = 0

        for row in rows[1:]:
            nombre = str(row[col_tienda]).strip() if col_tienda is not None and row[col_tienda] else None
            if not nombre or nombre == "None":
                continue

            # Buscar si ya existe (por nombre exacto)
            existing = Tienda.query.filter_by(nombre=nombre).first()

            data = {
                "co": str(row[col_co]).strip() if col_co is not None and row[col_co] else None,
                "ciudad": str(row[col_ciudad]).strip() if col_ciudad is not None and row[col_ciudad] else None,
                "concepto": str(row[col_concepto]).strip() if col_concepto is not None and row[col_concepto] else None,
                "correo_tienda": str(row[col_correo]).strip() if col_correo is not None and row[col_correo] else None,
                "direccion_tienda": str(row[col_direccion]).strip() if col_direccion is not None and row[col_direccion] else None,
            }

            if existing:
                for key, value in data.items():
                    setattr(existing, key, value)
                actualizados += 1
            else:
                tienda = Tienda(nombre=nombre, **data)
                db.session.add(tienda)
                importados += 1

        db.session.commit()

        print(f"\n--- Importación completada ---")
        print(f"Importados: {importados}")
        print(f"Actualizados: {actualizados}")
        print(f"Total de tiendas en BD: {Tienda.query.count()}")

        # Mostrar algunas
        tiendas = Tienda.query.limit(5).all()
        print("\nPrimeras 5 tiendas:")
        for t in tiendas:
            print(f"  {t.nombre} | CO: {t.co} | Ciudad: {t.ciudad} | Concepto: {t.concepto}")


if __name__ == "__main__":
    importar_tiendas()
